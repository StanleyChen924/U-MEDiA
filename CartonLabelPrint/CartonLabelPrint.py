import os
import sys
import json
import datetime
import configparser
import subprocess
import tkinter as tk
from tkinter import messagebox
import winsound
import importlib
import pymysql


class CartonSystemApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Carton Label 掃描控制系統")
        self.root.tk.call("tk", "scaling", 1.3)
        self.root.geometry("780x650")
        self.app_dir = (os.path.dirname(sys.executable) if getattr(sys, "frozen", False)
                        else os.path.dirname(os.path.abspath(__file__)))
        self.load_configs()
        self.log_dir = os.path.join(self.app_dir, "LOG")
        os.makedirs(self.log_dir, exist_ok=True)
        self.system_log_path = os.path.join(self.log_dir, f"system_log{datetime.datetime.now():%Y%m%d}.txt")
        self.s1_isn_list = []
        self.s2_isn_list = []
        self.work_order_set = set()
        self.carton_scanned_count = 0
        self.first_panel_id_of_carton = None
        self.load_label_data()
        self.create_widgets()
        self.root.protocol("WM_DELETE_WINDOW", self.close_application)

    def load_configs(self):
        self.config = configparser.ConfigParser()
        config_file = os.path.join(self.app_dir, "config.ini")
        if os.path.exists(config_file):
            self.config.read(config_file, encoding="utf-8")
        else:
            self.config["settings"] = {
                "csn_length": "18", "mac_length": "12", "Carton_ID": "CAR260900001",
                "excel_name": "2544259S1S2.xlsx", "label_name": "Barcode label.lab",
                "json_name": "label_data.json", "log_file": "scan_record.log",
                "tip_text": "請輸入或掃描資料（PANEL ID 23碼），並按 ENTER", "2D_AOI_Flag": "0"
            }
            with open(config_file, "w", encoding="utf-8") as f:
                self.config.write(f)
        self.mysql_config = configparser.ConfigParser()
        mysql_file = os.path.join(self.app_dir, "MySQLConfig.ini")
        if os.path.exists(mysql_file):
            self.mysql_config.read(mysql_file, encoding="utf-8")
        else:
            self.mysql_config["setting"] = {
                "MySQL_FLAG": "1", "MySQL_ServerIP": "10.4.5.13", "MySQL_username": "U94003",
                "MySQL_Password": "U94003", "MySQL_DB": "LITEON", "MySQL_PN": "2544259_S01+S2",
                "MySQL_LOT": "2544259-S01-003", "MySQL_QTY": "80", "MySQL_Carton": "1",
                "MySQL_Carton_Serial": "1"
            }
            with open(mysql_file, "w", encoding="utf-8") as f:
                self.mysql_config.write(f)
        get = lambda key, default: self.mysql_config.get("setting", key, fallback=default)
        self.cfg_pn = get("MySQL_PN", "2544259_S01+S2")
        self.cfg_dc = datetime.datetime.now().strftime("%Y%m%d")
        self.cfg_lot = get("MySQL_LOT", "2544259-S01-003")
        self.cfg_qty = get("MySQL_QTY", "80")
        self.cfg_carton = get("MySQL_Carton", "1")
        self.cfg_c_serial = get("MySQL_Carton_Serial", "1")
        self.last_panel_id = self.config.get("settings", "last_panel_id", fallback="")
        self.cfg_2d_aoi_flag = self.config.get("settings", "2D_AOI_Flag", fallback="0")

    def create_widgets(self):
        pad = {"padx": 10, "pady": 5}
        frame = tk.LabelFrame(self.root, text="參數設定 (可修改)")
        frame.pack(fill="x", **pad)
        labels = ["P/N:", "D/C:", "LOT(WO#):", "Q'ty:", "MySQL_Carton:", "MySQL_Carton_Serial:"]
        values = [self.cfg_pn, self.cfg_dc, self.cfg_lot, self.cfg_qty, self.cfg_carton, self.cfg_c_serial]
        self.entries = {}
        for i, label in enumerate(labels):
            row, col = divmod(i, 2)
            tk.Label(frame, text=label).grid(row=row, column=col * 2, sticky="e", padx=5, pady=2)
            entry = tk.Entry(frame, width=20)
            entry.insert(0, values[i])
            entry.grid(row=row, column=col * 2 + 1, sticky="w", padx=5, pady=2)
            self.entries[label] = entry
        scan = tk.LabelFrame(self.root, text="條碼掃描輸入區")
        scan.pack(fill="x", **pad)
        tk.Label(scan, text="Panel ID / Barcode:").pack(side="left", padx=5, pady=5)
        self.entry_barcode = tk.Entry(scan, width=40)
        self.entry_barcode.pack(side="left", fill="x", expand=True, padx=5, pady=5)
        self.entry_barcode.insert(0, self.last_panel_id)
        self.entry_barcode.bind("<Return>", self.process_barcode)
        status = tk.LabelFrame(self.root, text="系統提示與狀態")
        status.pack(fill="both", expand=True, **pad)
        self.lbl_tip = tk.Label(status, text=self.config.get("settings", "tip_text", fallback="請掃描"), fg="blue", anchor="w")
        self.lbl_tip.pack(fill="x", padx=5, pady=2)
        self.txt_log = tk.Text(status, height=10, state="disabled")
        self.txt_log.pack(fill="both", expand=True, padx=5, pady=5)
        buttons = tk.Frame(self.root)
        buttons.pack(fill="x", **pad)
        tk.Button(buttons, text="列印", command=self.manual_trigger_print, bg="green", fg="white", font=("Arial", 12, "bold")).pack(side="left", padx=5, pady=5)
        self.root.after_idle(self.entry_barcode.focus_set)

    def write_system_log(self, message):
        try:
            with open(self.system_log_path, "a", encoding="utf-8") as f:
                f.write(f"[{datetime.datetime.now():%Y-%m-%d %H:%M:%S}] {message}\n")
        except OSError:
            pass

    def log_message(self, message, color=None):
        self.txt_log.config(state="normal")
        text = f"[{datetime.datetime.now():%H:%M:%S}] {message}\n"
        if color:
            tag = f"log_{color}"
            self.txt_log.tag_configure(tag, foreground=color)
            self.txt_log.insert(tk.END, text, tag)
        else:
            self.txt_log.insert(tk.END, text)
        self.txt_log.config(state="disabled")
        self.txt_log.see(tk.END)
        self.write_system_log(message)

    def log_exception(self, error):
        self.write_system_log(f"{type(error).__name__}: {error}")

    def play_sound(self, name):
        path = os.path.join(self.app_dir, name)
        if not os.path.exists(path):
            self.log_message(f"⚠️ 找不到音效檔案: {path}")
            return
        try:
            winsound.PlaySound(path, winsound.SND_FILENAME | winsound.SND_NODEFAULT)
        except (RuntimeError, OSError) as e:
            self.log_message(f"⚠️ 音效播放失敗: {e}")

    def process_barcode(self, event=None):
        panel_id = self.entry_barcode.get().strip()
        self.entry_barcode.delete(0, tk.END)
        if not panel_id:
            return
        self.last_panel_id = panel_id
        self.config.set("settings", "last_panel_id", panel_id)
        with open(os.path.join(self.app_dir, "config.ini"), "w", encoding="utf-8") as f:
            self.config.write(f)
        self.log_message(f"讀取到 Panel ID: {panel_id}")
        pn = self.entries["P/N:"].get().strip()
        lot = self.entries["LOT(WO#):"].get().strip()
        qty = int(self.entries["Q'ty:"].get().strip())
        carton = self.entries["MySQL_Carton:"].get().strip()
        c_serial = int(self.entries["MySQL_Carton_Serial:"].get().strip())
        if c_serial == 1 and self.carton_scanned_count == 0:
            self.reset_label_progress()
        if self.first_panel_id_of_carton is None:
            self.first_panel_id_of_carton = panel_id
        excel_name = self.config.get("settings", "excel_name", fallback="2544259S1S2.xlsx").strip()
        excel_file = excel_name if os.path.isabs(excel_name) else os.path.join(self.app_dir, excel_name)
        if not os.path.exists(excel_file):
            messagebox.showerror("錯誤", f"找不到 Excel 檔案: {excel_name}\n搜尋路徑: {excel_file}")
            return
        try:
            from openpyxl import load_workbook
            from openpyxl.cell.cell import MergedCell
            wb = load_workbook(excel_file)
            ws = wb.active
            headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None and str(c.value).strip()}
            required = ("工單號", "併板主PNL", "序號(S1SN)", "序號(S2SN)")
            missing = [h for h in required if h not in headers]
            if missing:
                raise ValueError(f"Excel 缺少欄位: {', '.join(missing)}")
            panel_col, s1_col, s2_col = headers["併板主PNL"], headers["序號(S1SN)"], headers["序號(S2SN)"]
            start = max(s1_col, s2_col) + 1
            runtime_headers = ("Flag", "刷Barcode時間", "箱號", "流水號", "SPI_T", "SPI_T_StopTime", "SPI_B", "SPI_B_StopTime", "AOI_T", "AOI_T_StopTime", "AOI_B", "AOI_B_StopTime", "AOI_B_B", "AOI_B_B_StopTime", "AOI_B_T", "AOI_B_T_StopTime")
            runtime = {h: start + i for i, h in enumerate(runtime_headers)}
            for h, col in runtime.items():
                cell = ws.cell(1, col)
                if not isinstance(cell, MergedCell) and cell.value is None:
                    cell.value = h
            rows = [(n, ws[n]) for n in range(2, ws.max_row + 1) if str(ws.cell(n, panel_col).value or "").strip() == panel_id]
            if not rows:
                self.log_message(f"❌ Excel 中找不到該 Panel ID [{panel_id}] 的資料", "red")
                self.play_sound("buzz.wav")
                return
            successful = []
            skipped = 0
            for n, row in rows:
                serials = [(str(row[s1_col - 1].value or "").strip(), "S1"), (str(row[s2_col - 1].value or "").strip(), "S2")]
                serials = [(x, side) for x, side in serials if x]
                if str(row[runtime["Flag"] - 1].value or "").strip() == "Y":
                    skipped += 1
                    self.log_message(f"⚠️ PNL [{panel_id}] 已刷過主條碼 (Flag=Y)，跳過")
                    continue
                if not serials:
                    raise ValueError(f"PNL [{panel_id}] 沒有可驗證的 S1SN/S2SN")
                station_data = {}
                for isn, side in serials:
                    data = self.verify_database(isn)
                    if not data:
                        raise ValueError(f"{side} iSN [{isn}] 未通過資料庫(SPI/AOI)檢驗")
                    station_data[isn] = data
                successful.append((n, row, serials, station_data))
            if not successful and skipped == len(rows):
                self.play_sound("buzz.wav")
                self.log_message(f"❌ Panel ID [{panel_id}] 已全部刷過，拒絕重複掃描", "red")
                return
            for n, row, serials, station_data in successful:
                now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                ws.cell(n, runtime["Flag"]).value = "Y"
                ws.cell(n, runtime["刷Barcode時間"]).value = now
                ws.cell(n, runtime["箱號"]).value = carton
                ws.cell(n, runtime["流水號"]).value = c_serial
                for isn, side in serials:
                    target = self.s1_isn_list if side == "S1" else self.s2_isn_list
                    if isn not in target:
                        target.append(isn)
                    self.log_message(f"✅ {side} iSN [{isn}] 成功過站。流水號: {c_serial}")
                    c_serial += 1
                data = station_data[serials[0][0]]
                for station in ("SPI_T", "SPI_B", "AOI_T", "AOI_B", "AOI_B_B", "AOI_B_T"):
                    ws.cell(n, runtime[station]).value = data[station]
                    stop = data[f"{station}_StopTime"]
                    ws.cell(n, runtime[f"{station}_StopTime"]).value = stop.strftime("%Y-%m-%d %H:%M:%S") if stop else ""
                self.work_order_set.add(str(row[0].value or "").strip())
            self.play_sound("pass.wav")
            wb.save(excel_file)
            self.entries["MySQL_Carton_Serial:"].delete(0, tk.END)
            self.entries["MySQL_Carton_Serial:"].insert(0, str(c_serial))
            self.save_runtime_settings()
            self.write_label_json(pn, lot, qty)
            self.carton_scanned_count += sum(len(x[2]) for x in successful)
            self.log_message(f"本箱累計數量: {self.carton_scanned_count}/{qty}")
            if self.carton_scanned_count >= qty:
                self.trigger_print()
        except Exception as e:
            self.log_exception(e)
            self.play_sound("buzz.wav")
            messagebox.showerror("系統錯誤", f"處理過程中發生錯誤: {e}")

    def save_runtime_settings(self):
        for entry, key in (("P/N:", "MySQL_PN"), ("LOT(WO#):", "MySQL_LOT"), ("Q'ty:", "MySQL_QTY"), ("MySQL_Carton:", "MySQL_Carton"), ("MySQL_Carton_Serial:", "MySQL_Carton_Serial")):
            self.mysql_config.set("setting", key, self.entries[entry].get().strip())
        self.mysql_config.set("setting", "mysql_job", self.entries["LOT(WO#):"].get().strip())
        with open(os.path.join(self.app_dir, "MySQLConfig.ini"), "w", encoding="utf-8") as f:
            self.mysql_config.write(f)
        self.config.set("settings", "last_panel_id", self.last_panel_id)
        with open(os.path.join(self.app_dir, "config.ini"), "w", encoding="utf-8") as f:
            self.config.write(f)

    def close_application(self):
        self.save_runtime_settings()
        self.root.destroy()

    def load_label_data(self):
        path = os.path.join(self.app_dir, self.config.get("settings", "json_name", fallback="label_data.json"))
        if not os.path.exists(path):
            return
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            self.work_order_set = set(data.get("WorkOrders", []))
            self.s1_isn_list = [x for key in ("QRCode1", "QRCode3") for x in str(data.get(key, "")).split(",") if x.strip()]
            self.s2_isn_list = [x for key in ("QRCode5", "QRCode4") for x in str(data.get(key, "")).split(",") if x.strip()]
            self.carton_scanned_count = len(self.s1_isn_list) + len(self.s2_isn_list)
        except (OSError, json.JSONDecodeError) as e:
            self.log_exception(e)

    def verify_database(self, isn):
        try:
            db = importlib.import_module("pymysql")
            conn = db.connect(host=self.mysql_config.get("setting", "MySQL_ServerIP"), user=self.mysql_config.get("setting", "MySQL_username"), password=self.mysql_config.get("setting", "MySQL_Password"), database=self.mysql_config.get("setting", "MySQL_DB"), charset="utf8", cursorclass=db.cursors.DictCursor)
            cur = conn.cursor()
            stations = ("SPI_T", "SPI_B", "AOI_T", "AOI_B", "AOI_B_B", "AOI_B_T")
            fields = ", ".join(f"c.{s}, (SELECT errorCode FROM {s} WHERE iSN=c.iSN ORDER BY StopTime DESC, Serial DESC LIMIT 1) AS {s}_ErrorCode, (SELECT StopTime FROM {s} WHERE iSN=c.iSN ORDER BY StopTime DESC, Serial DESC LIMIT 1) AS {s}_StopTime" for s in stations)
            cur.execute(f"SELECT {fields} FROM CARD c WHERE c.iSN=%s LIMIT 1", (isn,))
            result = cur.fetchone()
            cur.close(); conn.close()
            stations = stations if self.cfg_2d_aoi_flag == "1" else stations[:4]
            if not result:
                return None
            if all(result.get(s) is not None and result[s] > 100 and str(result.get(f"{s}_ErrorCode") or "").strip().upper() in ("PASS", "RPASS") for s in stations):
                return result
            self.log_message(f"❌ iSN [{isn}] 未通過資料庫(SPI/AOI)檢驗", "red")
            return None
        except Exception as e:
            self.log_message(f"⚠️ 資料庫連線失敗或查詢錯誤: {e}")
            return None

    def _qr_data(self, pn, lot, qty, carton=None):
        values = self.s1_isn_list + self.s2_isn_list
        data = {"PN": pn, "DateCode": self.entries["D/C:"].get().strip(), "LOT WO": lot, "QTY": str(qty), "QRCode1": ",".join(values[:40]), "QRCode5": ",".join(values[40:80]), "QRCode3": ",".join(values[80:120]), "QRCode4": ",".join(values[120:160])}
        if carton is not None:
            data["Carton"] = str(carton)
        return data

    def write_label_json(self, pn, lot, qty):
        path = os.path.join(self.app_dir, self.config.get("settings", "json_name", fallback="label_data.json"))
        with open(path, "w", encoding="utf-8") as f:
            json.dump(self._qr_data(pn, lot, qty), f, ensure_ascii=False, indent=2)

    def save_json_to_archive(self, pn, lot, qty, carton):
        folder = os.path.join(self.app_dir, "JSON")
        os.makedirs(folder, exist_ok=True)
        name = f"label_data_{self.first_panel_id_of_carton}_{carton}.json"
        try:
            with open(os.path.join(folder, name), "w", encoding="utf-8") as f:
                json.dump(self._qr_data(pn, lot, qty, carton), f, ensure_ascii=False, indent=2)
            self.log_message(f"✅ JSON 備份已存至: {name}")
        except OSError as e:
            self.log_message(f"❌ JSON 備份失敗: {e}", "red")

    def manual_trigger_print(self):
        if not self.carton_scanned_count:
            messagebox.showwarning("警告", "目前箱內無掃描資料，無法列印")
            return
        self.save_json_to_archive(self.entries["P/N:"].get().strip(), self.entries["LOT(WO#):"].get().strip(), self.entries["Q'ty:"].get().strip(), self.entries["MySQL_Carton:"].get().strip())
        self.trigger_print()

    def trigger_print(self):
        self.log_message("達到目標數量 Q'ty，開始執行 test.bat 列印程序...")
        batch = os.path.join(self.app_dir, "test.bat")
        output_file = os.path.join(self.app_dir, "temp_dos.txt")
        if not os.path.exists(batch):
            self.log_message("❌ 找不到 test.bat 檔案，無法列印")
            return
        try:
            if os.path.exists(output_file):
                os.remove(output_file)
            result = subprocess.run(["cmd.exe", "/c", batch], cwd=self.app_dir, timeout=120, check=False)
            output = open(output_file, "rb").read().decode("utf-8", errors="replace") if os.path.exists(output_file) else ""
            if result.returncode != 0 or "PASS" not in output.upper() or "[OK]" not in output.upper():
                self.play_sound("buzz.wav")
                self.log_message("❌ FAIL：未確認列印已成功傳送至印表機佇列，維持目前箱號與流水號", "red")
                return
            self.log_message("🚀 PASS [OK]：已成功傳送至印表機佇列。")
            carton = int(self.entries["MySQL_Carton:"].get().strip()) + 1
            self.entries["MySQL_Carton:"].delete(0, tk.END); self.entries["MySQL_Carton:"].insert(0, str(carton))
            self.entries["MySQL_Carton_Serial:"].delete(0, tk.END); self.entries["MySQL_Carton_Serial:"].insert(0, "1")
            self.save_runtime_settings()
            self.s1_isn_list.clear(); self.s2_isn_list.clear(); self.work_order_set.clear()
            self.carton_scanned_count = 0
            self.first_panel_id_of_carton = None
        except subprocess.TimeoutExpired:
            self.play_sound("buzz.wav")
            self.log_message("❌ FAIL：列印程序逾時，維持目前箱號與流水號", "red")
        except Exception as e:
            self.log_message(f"❌ 列印執行失敗: {e}")

    def clear_label_data(self):
        path = os.path.join(self.app_dir, self.config.get("settings", "json_name", fallback="label_data.json"))
        if not os.path.exists(path):
            return
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            for key in ("QRCode1", "QRCode5", "QRCode3", "QRCode4"):
                data[key] = ""
            data.pop("WorkOrders", None)
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except (OSError, json.JSONDecodeError) as e:
            self.log_message(f"⚠️ 清除列印暫存資料失敗: {e}")

    def reset_label_progress(self):
        self.s1_isn_list.clear(); self.s2_isn_list.clear(); self.work_order_set.clear()
        self.carton_scanned_count = 0
        self.clear_label_data()
        self.log_message("新箱第一片掃描，已清空上一箱 JSON 與計數，從 1 開始")

    def manual_reset_count(self):
        self.reset_label_progress()
        self.first_panel_id_of_carton = None
        messagebox.showinfo("清空完成", "目前箱內計數已清空")


if __name__ == "__main__":
    root = tk.Tk()
    app = CartonSystemApp(root)
    root.mainloop()
