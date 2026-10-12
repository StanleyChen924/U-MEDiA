import configparser
from datetime import datetime
import pathlib
import re
import tkinter as tk
from tkinter import messagebox, ttk

import pandas as pd
import pymysql
from openpyxl import load_workbook

import sys

# 修正此處：判斷是 .py 執行還是 .exe 執行
if getattr(sys, 'frozen', False):
    # 點擊 EXE 時：改抓 sys.executable (即 EXE 檔案實際存放的目錄)
    BASE_DIR = pathlib.Path(sys.executable).resolve().parent
else:
    # 執行原始碼 .py 時：維持原本的邏輯
    BASE_DIR = pathlib.Path(__file__).resolve().parent

# 您原本的程式碼不需要變動
OUTPUT_DIR = BASE_DIR / "ASNFile"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# OUTPUT_DIR = BASE_DIR / "ASNFile"

TEMPLATE_PATTERN = "*shipping list.xlsx"
LAST_WORK_ORDER_FILE = BASE_DIR / "last_work_order.ini"

DB_COLUMNS = [
    "COMPLETE_TIME", "RELATION_ORDER", "DN_ITEM", "CUST_PART_NO", "PART_NO",
    "CUSTOMER_SN", "TEMP_SN", "CUSTOMER_SN2", "REEL_NO", "QTY", "PALLET_NO",
    "CARTON_NO", "STATUS", "WORK_ORDER", "MAC", "MODEL_NAME", "MFG_Date",
    "Shipping_Date", "PO",
]

ALIASES = {
    "MODELNAME": "MODEL_NAME",
    "MODEL_NAME": "MODEL_NAME",
    "SHIPPINGDATE": "Shipping_Date",
    "MFGDATE": "MFG_Date",
}


def read_db_config(config_file="MySQLConfig.ini"):
    config = configparser.ConfigParser()
    try:
        config.read(config_file, encoding="utf-8")
        return {
            "host": config.get("setting", "MySQL_ServerIP"),
            "user": config.get("setting", "MySQL_username"),
            "password": config.get("setting", "MySQL_Password"),
            "database": config.get("setting", "MySQL_DB"),
        }
    except Exception as exc:
        messagebox.showerror("配置錯誤", f"無法讀取設定檔 {config_file}\n原因: {exc}")
        raise


def load_last_work_order():
    if not LAST_WORK_ORDER_FILE.exists():
        return ""
    config = configparser.ConfigParser()
    try:
        config.read(LAST_WORK_ORDER_FILE, encoding="utf-8")
        return config.get("settings", "last_work_order", fallback="")
    except Exception:
        return ""


def save_last_work_order(work_order):
    if work_order:
        config = configparser.ConfigParser()
        config["settings"] = {"last_work_order": work_order}
        with LAST_WORK_ORDER_FILE.open("w", encoding="utf-8") as handle:
            config.write(handle)


def increment_pallet_no(base_str, increment_by):
    if not base_str:
        return ""
    match = re.match(r"(.*?)(\d+)$", str(base_str))
    if match:
        prefix, number = match.groups()
        return f"{prefix}{int(number) + increment_by:0{len(number)}d}"
    return f"{base_str}_{increment_by}" if increment_by else str(base_str)


def normalise_heading(value):
    if value is None:
        return ""
    text = re.sub(r"[^a-zA-Z0-9]", "", str(value)).upper()
    return ALIASES.get(text, str(value).strip())


def find_template():
    files = sorted(BASE_DIR.glob(TEMPLATE_PATTERN))
    if not files:
        files = sorted(pathlib.Path.cwd().glob(TEMPLATE_PATTERN))
    return files[0] if files else None


def read_template_definition():
    template = find_template()
    if not template:
        return {}, []

    workbook = load_workbook(template, read_only=True, data_only=True)
    configurations = {}
    choices = []
    try:
        for sheet in workbook.worksheets:
            if normalise_heading(sheet.cell(1, 1).value) != "MODEL_NAME":
                continue

            headings = {
                col: normalise_heading(sheet.cell(1, col).value)
                for col in range(2, sheet.max_column + 1)
                if sheet.cell(1, col).value not in (None, "")
            }
            if not any(heading in DB_COLUMNS for heading in headings.values()):
                continue

            for row in range(3, sheet.max_row + 1):
                model_name = sheet.cell(row, 1).value
                if model_name in (None, ""):
                    continue

                model_name = str(model_name).strip()
                values = {}
                required = set()
                for col, heading in headings.items():
                    if heading not in DB_COLUMNS:
                        continue
                    value = sheet.cell(row, col).value
                    if isinstance(value, str) and value.strip().lower() == "v":
                        required.add(heading)
                    elif value not in (None, ""):
                        values[heading] = value

                configurations[model_name] = {
                    "values": values,
                    "required": required,
                }
                choices.append(model_name)
            break

        # Keep model selection available for older templates without a product
        # configuration sheet.
        if not configurations:
            for sheet in workbook.worksheets:
                heading_row = None
                model_col = None
                for row_number, row in enumerate(sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 30),
                    values_only=True,
                ), 1):
                    for col_number, value in enumerate(row, 1):
                        if normalise_heading(value) == "MODEL_NAME":
                            heading_row = row_number
                            model_col = col_number
                            break
                    if model_col:
                        break
                if not model_col:
                    continue
                for row in range(max(heading_row + 1, 3), sheet.max_row + 1):
                    value = sheet.cell(row, model_col).value
                    if value not in (None, "") and str(value).strip().lower() != "v":
                        text = str(value).strip()
                        if text not in choices:
                            choices.append(text)
                if choices:
                    break
    finally:
        workbook.close()
    return configurations, choices


def execute_export():
    work_order = entry_wo.get().strip()
    save_last_work_order(work_order)
    model_name = combo_model.get().strip()
    total_qty_text = entry_total_qty.get().strip()
    qty_per_carton_text = entry_qty_per_carton.get().strip()
    cartons_per_pallet_text = entry_cartons_per_pallet.get().strip()
    input_po = entry_po.get().strip()
    input_complete_time = entry_complete_time.get().strip()
    input_relation_order = entry_relation_order.get().strip()
    input_dn_item = entry_cs_shipping_notice.get().strip()
    input_part_no = entry_product_part_no.get().strip()
    input_shipping_date = entry_shipping_date.get().strip()
    input_carton_start = entry_carton_start.get().strip()
    start_pallet = entry_pallet.get().strip()

    if not work_order:
        messagebox.showwarning("警告", "請輸入工單號碼！")
        return

    model_config = model_configurations.get(model_name, {})
    required = model_config.get("required", set())
    input_values = {
        "WORK_ORDER": work_order,
        "QTY": total_qty_text,
        "PO": input_po,
        "COMPLETE_TIME": input_complete_time,
        "RELATION_ORDER": input_relation_order,
        "DN_ITEM": input_dn_item,
        "PART_NO": input_part_no,
        "Shipping_Date": input_shipping_date,
        "CARTON_NO": input_carton_start,
    }
    missing_required = [
        column
        for column in required
        if column not in {"MAC", "MFG_Date"}
        and not (model_name if column == "MODEL_NAME" else input_values.get(column, ""))
    ]
    unsupported_required = [
        column
        for column in required
        if column not in {"MODEL_NAME", "MFG_Date"} and column not in input_values
    ]
    if unsupported_required:
        messagebox.showerror(
            "範本設定錯誤",
            "範本將下列欄位標記為必填，但程式沒有對應的輸入欄位："
            + "、".join(sorted(unsupported_required)),
        )
        return
    if missing_required:
        messagebox.showwarning(
            "必填欄位",
            "請填寫 Excel 中標記為 v 的欄位："
            + "、".join(sorted(missing_required)),
        )
        return

    try:
        total_qty = int(total_qty_text)
        if total_qty <= 0:
            raise ValueError
    except ValueError:
        messagebox.showwarning("輸入錯誤", "QTY 總數量必須為大於 0 的整數！")
        return

    try:
        qty_per_carton = int(qty_per_carton_text)
        if qty_per_carton <= 0:
            raise ValueError
    except ValueError:
        messagebox.showwarning("輸入錯誤", "每箱數量必須為大於 0 的整數！")
        return

    try:
        cartons_per_pallet = int(cartons_per_pallet_text)
        if cartons_per_pallet <= 0:
            raise ValueError
    except ValueError:
        messagebox.showwarning("輸入錯誤", "每棧板箱數必須為大於 0 的整數！")
        return

    connection = None
    try:
        label_status.config(text="狀態：正在查詢資料庫...", fg="blue")
        root.update_idletasks()
        db_config = read_db_config()
        select_fields = ", ".join(DB_COLUMNS)
        sql = f"""SELECT {select_fields} FROM ASN_Log
                  WHERE WORK_ORDER = %s
                    AND CARTON_NO IS NOT NULL
                    AND TRIM(CARTON_NO) <> ''"""
        params = [work_order]
        if input_carton_start:
            sql += " AND CARTON_NO >= %s"
            params.append(input_carton_start)
        if model_name:
            sql += " AND MODEL_NAME = %s"
            params.append(model_name)
        sql += " ORDER BY CARTON_NO ASC LIMIT %s"
        params.append(total_qty)

        connection = pymysql.connect(**db_config, charset="utf8")
        df = pd.read_sql(sql, connection, params=tuple(params))
        if df.empty:
            label_status.config(text="狀態：查無資料", fg="orange")
            messagebox.showinfo("提示", "找不到符合條件且 CARTON_NO 有值的資料。")
            return

        if len(df) < total_qty:
            messagebox.showwarning(
                "提醒",
                f"資料庫查到的筆數不足，預計 {total_qty} 筆，已依實際資料筆數匯出。",
            )

        overrides = {
            "PO": input_po,
            "COMPLETE_TIME": input_complete_time,
            "RELATION_ORDER": input_relation_order,
            "DN_ITEM": input_dn_item,
            "PART_NO": input_part_no,
            "Shipping_Date": input_shipping_date,
        }
        for column, value in overrides.items():
            if value:
                df[column] = value
        for column, value in model_config.get("values", {}).items():
            if column not in overrides and column not in {
                "CARTON_NO", "MAC", "MODEL_NAME", "QTY",
            }:
                df[column] = value
        if "MFG_Date" in required and df["MFG_Date"].fillna("").astype(str).str.strip().eq("").any():
            messagebox.showwarning(
                "必填欄位",
                "資料庫或範本設定中有資料缺少 MFG_Date，請確認資料後再匯出。",
            )
            return
        database_has_mac = (
            df["MAC"].notna()
            & df["MAC"].astype(str).str.strip().ne("")
        )
        df["MAC"] = df["MAC"].where(database_has_mac, "")
        if "MAC" in required and df["MAC"].astype(str).str.strip().eq("").any():
            messagebox.showwarning(
                "必填欄位",
                "資料庫中有資料缺少 MAC，請確認資料後再匯出。",
            )
            return
        if model_name:
            df["MODEL_NAME"] = model_name
        
        # --- 修正後的棧板號碼計算邏輯 ---
        if start_pallet:
            pallet_mapping = {}
            unique_cartons = df["CARTON_NO"].unique()
            for carton_index, carton_no in enumerate(unique_cartons):
                current_pallet_offset = carton_index // cartons_per_pallet
                pallet_mapping[carton_no] = increment_pallet_no(start_pallet, current_pallet_offset)
            df["PALLET_NO"] = df["CARTON_NO"].map(pallet_mapping)
        # ---------------------------------

        
        """
        df["PALLET_NO"] = [
            increment_pallet_no(start_pallet, index // cartons_per_pallet)
            if start_pallet else value
            for index, value in enumerate(df["PALLET_NO"])
        ]
        """

        output = df.reindex(columns=DB_COLUMNS).fillna("")
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = OUTPUT_DIR / f"{work_order}_{datetime.now():%Y%m%d}.xlsx"
        output.to_excel(output_path, index=False, engine="openpyxl")
        label_status.config(text="狀態：匯出成功！", fg="green")
        messagebox.showinfo(
            "成功",
            f"檔案匯出成功！\n檔案位置：{output_path}\n共 {len(output)} 筆資料。",
        )
    except pymysql.MySQLError as exc:
        label_status.config(text="狀態：資料庫錯誤", fg="red")
        messagebox.showerror("資料庫錯誤", f"操作失敗：\n{exc}")
    except Exception as exc:
        label_status.config(text="狀態：發生錯誤", fg="red")
        messagebox.showerror("錯誤", f"發生預期之外的錯誤：\n{exc}")
    finally:
        if connection:
            connection.close()


def focus_next(widget):
    widget.focus_set()


def input_row(label, default="", widget_type="entry", values=()):
    row = ttk.Frame(frame)
    row.pack(fill=tk.X, pady=3)
    tk.Label(row, text=label, font=("Microsoft JhengHei", 10), width=22, anchor="w").pack(side=tk.LEFT)
    if widget_type == "combo":
        widget = ttk.Combobox(row, values=values, font=("Microsoft JhengHei", 10))
    else:
        widget = ttk.Entry(row, font=("Microsoft Jheng Hei", 10))
    if default:
        widget.insert(0, default)
    widget.pack(side=tk.LEFT, fill=tk.X, expand=True)
    return widget


model_configurations, model_choices = read_template_definition()
root = tk.Tk()
root.title("ASN Log 資料匯出工具")
root.geometry("560x580")
root.resizable(False, False)
frame = ttk.Frame(root, padding="20")
frame.pack(fill=tk.BOTH, expand=True)
tk.Label(frame, text="工單資料 Excel 匯出系統", font=("Microsoft JhengHei", 14, "bold")).pack(pady=(0, 10))

combo_model = input_row("ModelName：", widget_type="combo", values=model_choices)
entry_wo = input_row("請輸入工單號碼：", load_last_work_order() or "WOTQ7553D")
entry_total_qty = input_row("QTY 總數量（抓取筆數）：", "")
entry_po = input_row("請輸入 PO 號碼：")
entry_qty_per_carton = input_row("每箱數量：", "")
entry_cartons_per_pallet = input_row("每棧板箱數：", "")
entry_carton_start = input_row("CARTON_NO 起始值：")
entry_complete_time = input_row("COMPLETE_TIME：")
entry_relation_order = input_row("RELATION_ORDER：")
entry_cs_shipping_notice = input_row("CS出貨通知：")
entry_product_part_no = input_row("成品料號：")
entry_shipping_date = input_row("Shipping_Date：")
entry_pallet = input_row("PALLET_NO 起始值：", "")

field_entries = {
    "WORK_ORDER": entry_wo,
    "QTY": entry_total_qty,
    "PO": entry_po,
    "COMPLETE_TIME": entry_complete_time,
    "RELATION_ORDER": entry_relation_order,
    "DN_ITEM": entry_cs_shipping_notice,
    "PART_NO": entry_product_part_no,
    "Shipping_Date": entry_shipping_date,
    "CARTON_NO": entry_carton_start,
    "PALLET_NO": entry_pallet,
}
last_model_fields = set()


def apply_model_configuration(event=None):
    global last_model_fields
    selected_model = event.widget.get().strip() if event else combo_model.get().strip()
    configuration = model_configurations.get(selected_model, {})
    values = configuration.get("values", {})
    required = configuration.get("required", set())
    fields_to_refresh = last_model_fields | set(values) | required

    for column in fields_to_refresh:
        widget = field_entries.get(column)
        if widget:
            widget.delete(0, tk.END)

    for column, value in values.items():
        widget = field_entries.get(column)
        if widget:
            widget.insert(0, str(value))

    last_model_fields = set(values) | required


combo_model.bind("<<ComboboxSelected>>", apply_model_configuration)
if model_choices:
    combo_model.current(0)
    apply_model_configuration()

label_status = tk.Label(frame, text="狀態：準備就緒", font=("Microsoft JhengHei", 9), fg="gray")
label_status.pack(pady=5)
ttk.Button(frame, text="開始查詢並匯出 Excel", command=execute_export).pack(fill=tk.X, ipady=5)
ttk.Button(frame, text="關閉程式", command=root.destroy).pack(fill=tk.X, ipady=5, pady=(5, 0))

entries = [
    combo_model, entry_wo, entry_total_qty, entry_po, entry_qty_per_carton,
    entry_cartons_per_pallet, entry_carton_start, entry_complete_time,
    entry_relation_order, entry_cs_shipping_notice, entry_product_part_no,
    entry_shipping_date, entry_pallet,
]
for current, next_widget in zip(entries, entries[1:]):
    current.bind("<Return>", lambda event, widget=next_widget: focus_next(widget))
entries[-1].bind("<Return>", lambda event: execute_export())
combo_model.focus_set()
root.mainloop()
